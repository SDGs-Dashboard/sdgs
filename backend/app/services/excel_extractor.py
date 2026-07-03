from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from ..database import execute, fetch_all, fetch_one, get_connection
from .audit import log_action
from .extraction_results import clear_report_results, create_review_placeholder, next_proposed_update_id, record_result, summarize_report_results
from .matcher import cleaned_match, confidence_label, exact_match, fuzzy_score, normalize_text, score_match
from .workbook_importer import get_current_dashboard_value


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


@dataclass
class SheetSnapshot:
    name: str
    rows: list[list[Any]]
    header_index: int
    header: list[Any]
    header_lookup: dict[str, int]
    text_blob: str


def _to_numeric(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _significant_tokens(value: Any) -> set[str]:
    stop_words = {
        "table",
        "figure",
        "and",
        "the",
        "for",
        "with",
        "from",
        "that",
        "this",
        "eicv",
        "nisr",
        "report",
        "thematic",
        "sex",
        "age",
        "province",
        "district",
        "residence",
        "area",
        "urban",
        "rural",
        "household",
        "households",
        "head",
        "heads",
        "population",
        "groups",
        "quintile",
        "consumption",
        "according",
        "compared",
        "girls",
        "boys",
        "male",
        "female",
        "total",
    }
    return {
        token
        for token in re.split(r"[^a-z0-9]+", normalize_text(value))
        if token and token not in stop_words and not token.isdigit()
    }


def _sheet_title(sheet: SheetSnapshot) -> str:
    if sheet.rows and sheet.rows[0]:
        first = sheet.rows[0][0]
        if normalize_text(first):
            return str(first)
    return sheet.name


def _mapping_reference_candidates(mapping: dict[str, Any], context: dict[str, Any]) -> list[str]:
    candidates: list[str] = []
    source_reference = str(context.get("source_table_reference") or "").strip()
    table_no = str(mapping.get("table_no") or "").strip()
    table_title = str(mapping.get("table_title") or "").strip()

    for value in [
        source_reference,
        f"{table_no}: {table_title}" if table_no and table_title else "",
        table_title,
    ]:
        normalized = str(value or "").strip()
        if normalized and normalized not in candidates:
            candidates.append(normalized)
    if not candidates and table_no:
        candidates.append(table_no)
    return candidates


def _expected_report_text(mapping: dict[str, Any], report: dict[str, Any]) -> str:
    return str(mapping.get("report_name") or mapping.get("report_family") or report.get("report_name") or "").strip()


def _expected_column_text(mapping: dict[str, Any], context: dict[str, Any], report: dict[str, Any]) -> str:
    target_year = mapping.get("report_year") or mapping.get("latest_year") or context.get("target_year") or report.get("publication_year")
    column_label = str(mapping.get("column_label") or "").strip()
    if column_label and target_year not in (None, ""):
        return f"{column_label} / {target_year}"
    if column_label:
        return column_label
    return str(target_year or "").strip()


def _best_label_match(cells: list[tuple[int, int, Any]], labels: list[str]) -> dict[str, Any] | None:
    best: dict[str, Any] | None = None
    for row_index, col_index, value in cells:
        if not normalize_text(value):
            continue
        for label in labels:
            if not normalize_text(label):
                continue
            match_type = None
            score = 0.0
            if exact_match(value, label):
                match_type = "exact"
                score = 1.0
            elif cleaned_match(value, label):
                match_type = "cleaned"
                score = 0.94
            else:
                fuzzy = fuzzy_score(value, label)
                if fuzzy >= 0.6:
                    match_type = "fuzzy"
                    score = fuzzy
            if match_type is None:
                continue
            candidate = {
                "row_index": row_index,
                "col_index": col_index,
                "value": str(value),
                "label": label,
                "match_type": match_type,
                "score": score,
            }
            if best is None or candidate["score"] > best["score"]:
                best = candidate
    return best


def _sheet_reference_score(sheet: SheetSnapshot, references: list[str]) -> tuple[float, str | None]:
    if not references:
        return 0.0, None
    title = _sheet_title(sheet)
    normalized_blob = normalize_text(sheet.text_blob)
    best_score = 0.0
    best_reference = None
    for reference in references:
        if exact_match(title, reference):
            return 1.0, reference
        if cleaned_match(title, reference):
            best_score = max(best_score, 0.95)
            best_reference = reference
        if normalize_text(reference) in normalized_blob:
            if 0.92 > best_score:
                best_score = 0.92
                best_reference = reference
        fuzzy = max(fuzzy_score(title, reference), fuzzy_score(sheet.name, reference))
        if fuzzy > best_score:
            best_score = fuzzy
            best_reference = reference
    return best_score, best_reference


def _nearby_numeric_values(sheet: SheetSnapshot, row_index: int, col_index: int) -> list[tuple[float, str]]:
    values: list[tuple[float, str]] = []
    for row_offset in range(0, 2):
        current_row_index = row_index + row_offset
        if current_row_index >= len(sheet.rows):
            continue
        row = sheet.rows[current_row_index]
        for col_offset in range(0, 3):
            current_col_index = col_index + col_offset
            if current_col_index >= len(row):
                continue
            numeric_value = _to_numeric(row[current_col_index])
            if numeric_value is not None:
                values.append((numeric_value, f"R{current_row_index + 1}C{current_col_index + 1}"))
    return values


def _validate_mapping(mapping: dict[str, Any], context: dict[str, Any], report: dict[str, Any]) -> tuple[str | None, str | None]:
    if not str(mapping.get("report_name") or mapping.get("report_family") or "").strip():
        return "missing report name", "missing_mapping"
    if not _mapping_reference_candidates(mapping, context):
        return "missing table number", "missing_mapping"
    if not _candidate_row_labels(mapping, context):
        return "missing row label", "missing_mapping"
    if not _expected_column_text(mapping, context, report):
        return "missing column/year", "missing_mapping"
    return None, None


def _load_workbook_snapshot(file_path: str | Path) -> list[SheetSnapshot]:
    workbook = load_workbook(file_path, read_only=False, data_only=True)
    snapshots: list[SheetSnapshot] = []
    for sheet in workbook.worksheets:
        merged_lookup: dict[tuple[int, int], Any] = {}
        for merged_range in sheet.merged_cells.ranges:
            top_left_value = sheet.cell(merged_range.min_row, merged_range.min_col).value
            for row_number in range(merged_range.min_row, merged_range.max_row + 1):
                for column_number in range(merged_range.min_col, merged_range.max_col + 1):
                    merged_lookup[(row_number, column_number)] = top_left_value

        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        rows: list[list[Any]] = []
        for row_number in range(1, max_row + 1):
            if sheet.row_dimensions[row_number].hidden:
                continue
            row: list[Any] = []
            for column_number in range(1, max_col + 1):
                value = sheet.cell(row_number, column_number).value
                if value in (None, "") and (row_number, column_number) in merged_lookup:
                    value = merged_lookup[(row_number, column_number)]
                row.append(value)
            if any(normalize_text(cell) for cell in row):
                rows.append(row)
        header_index = 0
        header = rows[0] if rows else []
        for index, row in enumerate(rows[:10]):
            non_blank = sum(1 for cell in row if normalize_text(cell))
            if non_blank >= 2:
                header_index = index
                header = row
                break
        header_lookup = {normalize_text(value): position for position, value in enumerate(header) if normalize_text(value)}
        text_blob = "\n".join(" | ".join("" if cell is None else str(cell) for cell in row) for row in rows[:100])
        snapshots.append(
            SheetSnapshot(
                name=sheet.title,
                rows=rows,
                header_index=header_index,
                header=header,
                header_lookup=header_lookup,
                text_blob=text_blob,
            )
        )
    return snapshots


def _report_matches_mapping(report: dict[str, Any], mapping: dict[str, Any]) -> bool:
    report_type = normalize_text(report.get("report_type"))
    mapping_file_type = normalize_text(mapping.get("file_type"))
    report_family = normalize_text(report.get("report_family"))
    mapping_family = normalize_text(mapping.get("report_family"))
    mapping_report_name = normalize_text(mapping.get("report_name"))
    report_name = normalize_text(report.get("report_name"))

    if mapping_file_type and "excel" not in mapping_file_type and "pdf+excel" not in mapping_file_type:
        return False
    if mapping_family and report_family and mapping_family not in report_family and report_family not in mapping_family:
        return False
    if mapping_report_name and report_name and mapping_report_name not in report_name and report_name not in mapping_report_name:
        return False
    return report_type in {"excel", "xls", "xlsx", "csv"} or report_type == "excel"


def _get_mapping_context(mapping: dict[str, Any], report: dict[str, Any]) -> dict[str, Any]:
    target_year = mapping.get("report_year") or mapping.get("latest_year") or report.get("publication_year")
    params = (mapping["indicator"], mapping.get("series_code"), int(target_year)) if target_year not in (None, "") else None

    dashboard_row = None
    if params:
        dashboard_row = fetch_one(
            """
            SELECT *
            FROM dashboard_data
            WHERE indicator = ?
              AND COALESCE(series_code, '') = COALESCE(?, '')
              AND year = ?
            ORDER BY CASE WHEN ref_area = 'RW' THEN 0 ELSE 1 END, id
            LIMIT 1
            """,
            params,
        )
    if dashboard_row is None:
        dashboard_row = fetch_one(
            """
            SELECT *
            FROM dashboard_data
            WHERE indicator = ?
              AND COALESCE(series_code, '') = COALESCE(?, '')
            ORDER BY CASE WHEN ref_area = 'RW' THEN 0 ELSE 1 END, year DESC, id
            LIMIT 1
            """,
            (mapping["indicator"], mapping.get("series_code")),
        )

    indicator_row = fetch_one(
        """
        SELECT *
        FROM indicators
        WHERE indicator_code = ?
          AND COALESCE(series_code, '') = COALESCE(?, '')
        LIMIT 1
        """,
        (mapping["indicator"], mapping.get("series_code")),
    )

    context: dict[str, Any] = {}
    if indicator_row:
        context.update(dict(indicator_row))
    if dashboard_row:
        context.update(dict(dashboard_row))
    context["target_year"] = int(target_year) if target_year not in (None, "") else None
    return context


def _insert_extracted_table(report_id: str, sheet: SheetSnapshot) -> None:
    preview_rows = sheet.rows[:8]
    execute(
        """
        INSERT INTO extracted_tables (
            report_id, sheet_name, page_number, table_number, table_title,
            header_row_index, preview_json, text_snapshot, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            report_id,
            sheet.name,
            None,
            None,
            sheet.name,
            sheet.header_index,
            str(preview_rows),
            sheet.text_blob[:4000],
            utc_now(),
        ),
    )


def _direct_series_year_match(sheet: SheetSnapshot, mapping: dict[str, Any], report: dict[str, Any]) -> dict[str, Any] | None:
    normalized_header = {normalize_text(value): idx for idx, value in enumerate(sheet.header)}
    if "series_code" not in normalized_header and "indicator" not in normalized_header:
        return None

    series_col = normalized_header.get("series_code")
    indicator_col = normalized_header.get("indicator")
    target_series_code = str(mapping.get("series_code") or "").strip()
    target_indicator = str(mapping.get("indicator") or "").strip()
    target_year = mapping.get("report_year") or mapping.get("latest_year") or report.get("publication_year")

    year_lookup = {int(value): idx for idx, value in enumerate(sheet.header) if isinstance(value, int)}
    if not year_lookup:
        return None

    if target_year is None:
        target_year = max(year_lookup)
    target_year = int(target_year)
    year_col = year_lookup.get(target_year) or year_lookup.get(max(year_lookup))
    chosen_year = target_year if target_year in year_lookup else max(year_lookup)

    for row_index in range(sheet.header_index + 1, len(sheet.rows)):
        row = sheet.rows[row_index]
        series_match = bool(target_series_code and series_col is not None and exact_match(row[series_col], target_series_code))
        indicator_match = bool(target_indicator and indicator_col is not None and exact_match(row[indicator_col], target_indicator))
        if not (series_match or indicator_match):
            continue

        value = row[year_col] if year_col is not None and year_col < len(row) else None
        try:
            numeric_value = float(value) if value not in (None, "") else None
        except (TypeError, ValueError):
            numeric_value = None
        if numeric_value is None:
            continue

        score, note = score_match(
            table_number_match=False,
            table_title_match=False,
            row_label_match=True,
            column_label_match=True,
            fuzzy_only=False,
            series_code_match=series_match,
            indicator_code_match=indicator_match,
        )
        return {
            "year": chosen_year,
            "new_value": numeric_value,
            "table_or_sheet": sheet.name,
            "evidence_page": sheet.name,
            "source_evidence": f"{sheet.name}!R{row_index + 1}C{year_col + 1}",
            "matched_cell": f"R{row_index + 1}C{year_col + 1}",
            "confidence_score": score,
            "confidence_label": confidence_label(score),
            "extraction_method": "Excel direct row/year",
            "extraction_note": note,
        }
    return None


def _find_row_label_match(sheet: SheetSnapshot, label: str) -> tuple[int, int] | None:
    target = normalize_text(label)
    if not target:
        return None
    for row_index, row in enumerate(sheet.rows):
        for col_index, cell in enumerate(row):
            if exact_match(cell, target):
                return row_index, col_index
    return None


def _find_column_match(sheet: SheetSnapshot, label: str, year: int | None) -> tuple[int, int] | None:
    candidates: list[str] = []
    if label:
      candidates.append(label)
    if year is not None:
      candidates.append(str(year))

    for row_index in range(min(8, len(sheet.rows))):
        row = sheet.rows[row_index]
        for col_index, cell in enumerate(row):
            for candidate in candidates:
                if exact_match(cell, candidate):
                    return row_index, col_index
    return None


def _candidate_row_labels(mapping: dict[str, Any], context: dict[str, Any]) -> list[str]:
    labels: list[str] = []
    for value in [
        mapping.get("row_label"),
        context.get("district"),
        context.get("province"),
        context.get("urbanization"),
    ]:
        normalized = str(value or "").strip()
        if normalized and normalized not in labels:
            labels.append(normalized)

    ref_area = str(context.get("ref_area") or context.get("geography") or "").strip().upper()
    if ref_area in {"RW", "RWA"} or str(context.get("row_label") or "").strip().lower() == "rwanda":
        for label in ["All Rwanda", "Rwanda", "National", "Total"]:
            if label not in labels:
                labels.append(label)
    return labels


def _header_signatures(sheet: SheetSnapshot, depth: int = 5) -> dict[int, str]:
    preview_rows = sheet.rows[: max(depth, sheet.header_index + 2)]
    column_count = max((len(row) for row in preview_rows), default=0)
    signatures: dict[int, str] = {}
    for col_index in range(column_count):
        parts: list[str] = []
        for row in preview_rows:
            if col_index >= len(row):
                continue
            token = normalize_text(row[col_index])
            if token and token not in parts:
                parts.append(token)
        signatures[col_index] = " ".join(parts)
    return signatures


def _candidate_column_terms(mapping: dict[str, Any], context: dict[str, Any], target_year: int | None) -> list[str]:
    terms: list[str] = []
    column_label = str(mapping.get("column_label") or "").strip()
    if column_label:
        for chunk in column_label.replace("/", "|").split("|"):
            token = chunk.strip()
            if token:
                terms.append(token)

    text_sources = [
        mapping.get("table_title"),
        mapping.get("report_indicator_name"),
        mapping.get("series"),
        mapping.get("dashboard_description"),
        context.get("description"),
        context.get("dashboard_description"),
        context.get("education"),
    ]
    phrases = [
        "net attendance rate",
        "gross attendance rate",
        "gender parity index",
        "primary",
        "secondary",
        "headcount poverty rate",
        "poverty rate",
        "extreme poverty",
        "multidimensional poverty index",
        "incidence",
        "intensity",
    ]
    combined_text = " ".join(normalize_text(value) for value in text_sources if value)
    for phrase in phrases:
        if phrase in combined_text and phrase not in terms:
            terms.append(phrase)

    if target_year is not None:
        terms.append(str(target_year))
    return terms


def _semantic_table_score(sheet: SheetSnapshot, mapping: dict[str, Any], context: dict[str, Any], report: dict[str, Any]) -> tuple[float, str]:
    title = _sheet_title(sheet)
    mapping_fields = _mapping_reference_candidates(mapping, context)
    if not mapping_fields:
        return 0.0, "no mapped table reference"

    best_score = 0.0
    best_reason = ""
    title_tokens = _significant_tokens(title)
    report_tokens = _significant_tokens(report.get("report_name"))

    for field in mapping_fields:
        if not normalize_text(field):
            continue
        ratio = fuzzy_score(title, field)
        overlap = len(title_tokens & _significant_tokens(field))
        overlap_ratio = overlap / max(len(_significant_tokens(field)), 1)
        score = max(ratio, overlap_ratio)
        if overlap >= 2:
            score = max(score, min(0.9, 0.45 + 0.12 * overlap))
        if report_tokens and report_tokens & _significant_tokens(field):
            score += 0.05
        if score > best_score:
            best_score = score
            best_reason = f"title similarity to `{str(field)[:80]}`"

    return min(best_score, 0.95), best_reason or "semantic title similarity"


def _best_column_match(sheet: SheetSnapshot, terms: list[str], target_year: int | None) -> tuple[int | None, float]:
    signatures = _header_signatures(sheet)
    best_column = None
    best_score = 0.0
    normalized_terms = [normalize_text(term) for term in terms if normalize_text(term)]

    for col_index, signature in signatures.items():
        score = 0.0
        for term in normalized_terms:
            if term and term in signature:
                score += 2.0 if " " in term else 1.0
        if target_year is not None and str(target_year) in signature:
            score += 2.0
        if score > best_score:
            best_score = score
            best_column = col_index

    return (best_column if best_score > 0 else None, best_score)


def _best_column_index(sheet: SheetSnapshot, terms: list[str], target_year: int | None) -> int | None:
    return _best_column_match(sheet, terms, target_year)[0]


def _extract_value_from_row(row: list[Any], row_index: int, col_index: int | None) -> tuple[float | None, str | None]:
    if col_index is not None and col_index < len(row):
        numeric_value = _to_numeric(row[col_index])
        if numeric_value is not None:
            return numeric_value, f"R{row_index + 1}C{col_index + 1}"

    for candidate_index, cell in enumerate(row[1:], start=1):
        numeric_value = _to_numeric(cell)
        if numeric_value is not None:
            return numeric_value, f"R{row_index + 1}C{candidate_index + 1}"
    return None, None


def _semantic_table_suggestion(sheet: SheetSnapshot, mapping: dict[str, Any], report: dict[str, Any], context: dict[str, Any]) -> dict[str, Any] | None:
    if normalize_text(sheet.name) == "table of contents" or normalize_text(sheet.name).startswith("figure"):
        return None

    reference_candidates = _mapping_reference_candidates(mapping, context)
    if not reference_candidates:
        return None

    score, reason = _semantic_table_score(sheet, mapping, context, report)
    if score < 0.63:
        return None

    table_no = str(mapping.get("table_no") or "").strip()
    normalized_sheet_title = normalize_text(_sheet_title(sheet))
    if table_no:
        if normalize_text(table_no) not in normalize_text(sheet.name) and normalize_text(table_no) not in normalized_sheet_title:
            return None
        if max((fuzzy_score(_sheet_title(sheet), reference) for reference in reference_candidates), default=0.0) < 0.82:
            return None
    elif max((fuzzy_score(_sheet_title(sheet), reference) for reference in reference_candidates), default=0.0) < 0.82:
        return None

    target_year = context.get("target_year")
    row_match = None
    matched_label = None
    for label in _candidate_row_labels(mapping, context):
        row_match = _find_row_label_match(sheet, label)
        if row_match:
            matched_label = label
            break

    matched_cell = None
    numeric_value = None
    if row_match:
        row_index, _ = row_match
        col_index, col_score = _best_column_match(sheet, _candidate_column_terms(mapping, context, target_year), target_year)
        generic_row = normalize_text(matched_label) in {"rwanda", "all rwanda", "national", "total"}
        if col_index is not None and col_score >= 2.0 and not generic_row:
            numeric_value, matched_cell = _extract_value_from_row(sheet.rows[row_index], row_index, col_index)

    confidence = min(0.82, score if numeric_value is None else score + 0.08)
    return {
        "year": target_year or report.get("publication_year") or datetime.now().year,
        "new_value": numeric_value,
        "table_or_sheet": sheet.name,
        "evidence_page": sheet.name,
        "source_evidence": f"{sheet.name}!{matched_cell}" if matched_cell else _sheet_title(sheet),
        "matched_cell": matched_cell,
        "confidence_score": confidence,
        "confidence_label": confidence_label(confidence),
        "extraction_method": "Related table suggestion",
        "extraction_note": f"{reason}; workbook mapping incomplete so review is required",
    }


def _metadata_match(sheet: SheetSnapshot, mapping: dict[str, Any], report: dict[str, Any], context: dict[str, Any]) -> dict[str, Any] | None:
    table_no = str(mapping.get("table_no") or "").strip()
    table_title = str(mapping.get("table_title") or "").strip()
    reference_candidates = _mapping_reference_candidates(mapping, context)
    target_year = mapping.get("report_year") or mapping.get("latest_year") or report.get("publication_year")
    target_year = int(target_year) if target_year not in (None, "") else None

    sheet_match = str(mapping.get("sheet_or_page") or "").strip()
    if sheet_match and not exact_match(sheet.name, sheet_match):
        return None

    normalized_blob = normalize_text(sheet.text_blob)
    table_number_match = bool(table_no and (normalize_text(table_no) in normalize_text(sheet.name) or normalize_text(table_no) in normalized_blob))
    table_title_match = bool(table_title and (normalize_text(table_title) in normalized_blob or fuzzy_score(sheet.rows[0][0] if sheet.rows and sheet.rows[0] else "", table_title) >= 0.72))
    reference_match = any(
        normalize_text(reference) in normalized_blob or fuzzy_score(_sheet_title(sheet), reference) >= 0.72
        for reference in reference_candidates
    )

    if reference_candidates:
        if not reference_match and not sheet_match:
            return None
    elif table_title:
        if not table_title_match and not sheet_match:
            return None
    elif table_no:
        if not table_number_match and not sheet_match:
            return None

    row_match = None
    for label in _candidate_row_labels(mapping, context):
        row_match = _find_row_label_match(sheet, label)
        if row_match:
            break

    column_match = None
    column_label = str(mapping.get("column_label") or (target_year if target_year else "")).strip()
    if column_label:
        column_match = _find_column_match(sheet, column_label, target_year)
    if column_match is None:
        candidate_column = _best_column_index(sheet, _candidate_column_terms(mapping, context, target_year), target_year)
        if candidate_column is not None:
            column_match = (sheet.header_index, candidate_column)

    if row_match and column_match:
        row_index, _ = row_match
        _, col_index = column_match
        row = sheet.rows[row_index]
        numeric_value, matched_cell = _extract_value_from_row(row, row_index, col_index)
        if numeric_value is not None:
            score, note = score_match(
                table_number_match=table_number_match,
                table_title_match=table_title_match,
                row_label_match=True,
                column_label_match=True,
                fuzzy_only=False,
            )
            return {
                "year": target_year or int(column_label),
                "new_value": numeric_value,
                "table_or_sheet": sheet.name,
                "evidence_page": sheet.name,
                "source_evidence": f"{sheet.name}!{matched_cell}",
                "matched_cell": matched_cell,
                "confidence_score": score,
                "confidence_label": confidence_label(score),
                "extraction_method": "Excel metadata match",
                "extraction_note": note,
            }

    row_label = str(mapping.get("row_label") or "").strip()
    if row_label and column_match:
        best_ratio = 0.0
        best_row_index = None
        best_row = None
        for row_index, row in enumerate(sheet.rows[sheet.header_index + 1 :], start=sheet.header_index + 1):
            left_value = next((cell for cell in row if normalize_text(cell)), None)
            ratio = fuzzy_score(left_value, row_label)
            if ratio > best_ratio:
                best_ratio = ratio
                best_row_index = row_index
                best_row = row
        if best_ratio >= 0.86 and best_row_index is not None and best_row is not None:
            col_index = column_match[1]
            numeric_value, matched_cell = _extract_value_from_row(best_row, best_row_index, col_index)
            if numeric_value is not None:
                score, note = score_match(
                    table_number_match=table_number_match,
                    table_title_match=table_title_match,
                    row_label_match=False,
                    column_label_match=True,
                    fuzzy_only=True,
                )
                return {
                    "year": target_year or int(column_label),
                    "new_value": numeric_value,
                    "table_or_sheet": sheet.name,
                    "evidence_page": sheet.name,
                    "source_evidence": f"{sheet.name}!{matched_cell}",
                    "matched_cell": matched_cell,
                    "confidence_score": score,
                    "confidence_label": confidence_label(score),
                    "extraction_method": "Fuzzy suggestion",
                    "extraction_note": f"{note}; fuzzy row label similarity {best_ratio:.2f}",
                }
    return None


def extract_report(report_id: str) -> dict[str, Any]:
    report = fetch_one("SELECT * FROM reports WHERE report_id = ?", (report_id,))
    if not report:
        raise ValueError("Report not found.")

    snapshots = _load_workbook_snapshot(report["file_path"])
    mappings = fetch_all("SELECT * FROM source_mapping ORDER BY mapping_id")

    clear_report_results(report_id)

    for sheet in snapshots:
        _insert_extracted_table(report_id, sheet)

    proposals = 0
    for mapping in mappings:
        context = _get_mapping_context(mapping, report)
        expected_report = _expected_report_text(mapping, report)
        expected_table = " | ".join(_mapping_reference_candidates(mapping, context))
        expected_row = " | ".join(_candidate_row_labels(mapping, context))
        expected_column = _expected_column_text(mapping, context, report)

        try:
            if not _report_matches_mapping(report, mapping):
                record_result(
                    report_id=report_id,
                    mapping=mapping,
                    status="missing_mapping",
                    reason="report does not match uploaded file",
                    expected_report=expected_report,
                    expected_table=expected_table,
                    expected_row=expected_row,
                    expected_column=expected_column,
                    debug_message="Mapping row belongs to a different report family, report name, or file type.",
                )
                continue

            validation_reason, validation_status = _validate_mapping(mapping, context, report)
            if validation_reason:
                record_result(
                    report_id=report_id,
                    mapping=mapping,
                    status=validation_status or "missing_mapping",
                    reason=validation_reason,
                    expected_report=expected_report,
                    expected_table=expected_table,
                    expected_row=expected_row,
                    expected_column=expected_column,
                )
                continue

            sheet_rankings: list[tuple[float, SheetSnapshot]] = []
            for sheet in snapshots:
                sheet_score, _ = _sheet_reference_score(sheet, _mapping_reference_candidates(mapping, context))
                if sheet_score >= 0.6:
                    sheet_rankings.append((sheet_score, sheet))

            sheet_rankings.sort(key=lambda item: item[0], reverse=True)
            if not sheet_rankings:
                record_result(
                    report_id=report_id,
                    mapping=mapping,
                    status="not_found",
                    reason="table not found",
                    expected_report=expected_report,
                    expected_table=expected_table,
                    expected_row=expected_row,
                    expected_column=expected_column,
                )
                continue

            if len(sheet_rankings) > 1 and sheet_rankings[0][0] >= 0.8 and abs(sheet_rankings[0][0] - sheet_rankings[1][0]) <= 0.03:
                placeholder_id = create_review_placeholder(
                    report_id=report_id,
                    report=report,
                    mapping=mapping,
                    year=context.get("target_year") or report.get("publication_year"),
                    old_value=get_current_dashboard_value(
                        mapping["indicator"],
                        mapping.get("series_code"),
                        int(context.get("target_year") or report.get("publication_year") or datetime.now().year),
                    ),
                    status="Needs Review",
                    confidence_score=sheet_rankings[0][0],
                    confidence_label=confidence_label(sheet_rankings[0][0]),
                    extraction_method="Excel ambiguous sheet match",
                    extraction_note="Multiple possible sheets matched this mapping. Enter the confirmed value after manual review.",
                    table_or_sheet=_sheet_title(sheet_rankings[0][1]),
                    evidence_page=sheet_rankings[0][1].name,
                    source_evidence=sheet_rankings[0][1].name,
                )
                record_result(
                    report_id=report_id,
                    mapping=mapping,
                    status="ambiguous_match",
                    reason="multiple possible matches found",
                    expected_report=expected_report,
                    expected_table=expected_table,
                    expected_row=expected_row,
                    expected_column=expected_column,
                    confidence_score=sheet_rankings[0][0],
                    source_sheet_page=sheet_rankings[0][1].name,
                    matched_table=_sheet_title(sheet_rankings[0][1]),
                    proposed_update_id=placeholder_id,
                )
                proposals += 1
                continue

            best_sheet = sheet_rankings[0][1]
            candidate = _direct_series_year_match(best_sheet, mapping, report) or _metadata_match(best_sheet, mapping, report, context)

            if candidate is None:
                all_cells = [
                    (row_index, col_index, cell)
                    for row_index, row in enumerate(best_sheet.rows)
                    for col_index, cell in enumerate(row)
                ]
                row_match = _best_label_match(all_cells, _candidate_row_labels(mapping, context))
                if row_match is None:
                    record_result(
                        report_id=report_id,
                        mapping=mapping,
                        status="not_found",
                        reason="row label not found",
                        expected_report=expected_report,
                        expected_table=expected_table,
                        expected_row=expected_row,
                        expected_column=expected_column,
                        source_sheet_page=best_sheet.name,
                        matched_table=_sheet_title(best_sheet),
                    )
                    continue

                header_cells = [
                    (row_index, col_index, cell)
                    for row_index, row in enumerate(best_sheet.rows[: min(len(best_sheet.rows), 12)])
                    for col_index, cell in enumerate(row)
                ]
                column_match = _best_label_match(header_cells, [token for token in _candidate_column_terms(mapping, context, context.get("target_year")) if token])
                if column_match is None:
                    record_result(
                        report_id=report_id,
                        mapping=mapping,
                        status="not_found",
                        reason="year/column not found",
                        expected_report=expected_report,
                        expected_table=expected_table,
                        expected_row=expected_row,
                        expected_column=expected_column,
                        closest_matched_row=row_match["value"],
                        confidence_score=row_match["score"],
                        source_sheet_page=best_sheet.name,
                        matched_table=_sheet_title(best_sheet),
                    )
                    continue

                nearby_values = _nearby_numeric_values(best_sheet, row_match["row_index"], column_match["col_index"])
                if not nearby_values:
                    record_result(
                        report_id=report_id,
                        mapping=mapping,
                        status="not_found",
                        reason="value empty",
                        expected_report=expected_report,
                        expected_table=expected_table,
                        expected_row=expected_row,
                        expected_column=expected_column,
                        closest_matched_row=row_match["value"],
                        closest_matched_column=column_match["value"],
                        confidence_score=min(row_match["score"], column_match["score"]),
                        source_sheet_page=best_sheet.name,
                        matched_table=_sheet_title(best_sheet),
                    )
                    continue

                distinct_values = {value for value, _ in nearby_values}
                if len(distinct_values) > 1:
                    placeholder_id = create_review_placeholder(
                        report_id=report_id,
                        report=report,
                        mapping=mapping,
                        year=context.get("target_year") or report.get("publication_year"),
                        old_value=get_current_dashboard_value(
                            mapping["indicator"],
                            mapping.get("series_code"),
                            int(context.get("target_year") or report.get("publication_year") or datetime.now().year),
                        ),
                        status="Needs Review",
                        confidence_score=min(row_match["score"], column_match["score"]),
                        confidence_label=confidence_label(min(row_match["score"], column_match["score"])),
                        extraction_method="Excel conflicting nearby values",
                        extraction_note="More than one nearby numeric value matched this row/column. Enter the confirmed value after manual review.",
                        table_or_sheet=best_sheet.name,
                        evidence_page=best_sheet.name,
                        source_evidence=best_sheet.name,
                    )
                    record_result(
                        report_id=report_id,
                        mapping=mapping,
                        status="ambiguous_match",
                        reason="multiple possible matches found",
                        expected_report=expected_report,
                        expected_table=expected_table,
                        expected_row=expected_row,
                        expected_column=expected_column,
                        closest_matched_row=row_match["value"],
                        closest_matched_column=column_match["value"],
                        confidence_score=min(row_match["score"], column_match["score"]),
                        source_sheet_page=best_sheet.name,
                        matched_table=_sheet_title(best_sheet),
                        debug_message="Nearby numeric values conflict around the matched row/column.",
                        proposed_update_id=placeholder_id,
                    )
                    proposals += 1
                    continue

                extracted_value, matched_cell = nearby_values[0]
                confidence = min(row_match["score"], column_match["score"], sheet_rankings[0][0])
                match_reason = "extracted"
                result_status = "extracted"
                if row_match["match_type"] == "fuzzy" or column_match["match_type"] == "fuzzy":
                    if confidence < 0.8:
                        result_status = "ambiguous_match"
                        match_reason = "multiple possible matches found"
                    else:
                        match_reason = "cleaned text match"
                proposal_status = "Needs Review" if result_status == "ambiguous_match" else ("Pending Review" if confidence < 0.9 else "Ready")
                year = int(context.get("target_year") or report.get("publication_year") or datetime.now().year)
                old_value = get_current_dashboard_value(mapping["indicator"], mapping.get("series_code"), year)
                difference = None if old_value is None else extracted_value - old_value
                update_id = next_proposed_update_id(report_id)
                execute(
                    """
                    INSERT INTO proposed_updates (
                        update_id, mapping_id, indicator, series_code, year, old_value, new_value,
                        difference, unit_code, source_report, source_report_id, table_or_sheet, evidence_page,
                        extraction_date, status, reviewer_comment, confidence_score, confidence_label,
                        extraction_method, extraction_note, source_evidence, matched_cell, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        update_id,
                        mapping["mapping_id"],
                        mapping["indicator"],
                        mapping.get("series_code"),
                        year,
                        old_value,
                        extracted_value,
                        difference,
                        mapping.get("unit_code"),
                        report["report_name"],
                        report_id,
                        best_sheet.name,
                        best_sheet.name,
                        utc_now(),
                        proposal_status,
                        None,
                        confidence,
                        confidence_label(confidence),
                        "Excel nearby cell search",
                        match_reason,
                        f"{best_sheet.name}!{matched_cell}",
                        matched_cell,
                        utc_now(),
                        utc_now(),
                    ),
                )
                proposals += 1
                record_result(
                    report_id=report_id,
                    mapping=mapping,
                    status=result_status,
                    reason=match_reason,
                    expected_report=expected_report,
                    expected_table=expected_table,
                    expected_row=expected_row,
                    expected_column=expected_column,
                    closest_matched_row=row_match["value"],
                    closest_matched_column=column_match["value"],
                    confidence_score=confidence,
                    source_sheet_page=best_sheet.name,
                    extracted_value=extracted_value,
                    matched_table=_sheet_title(best_sheet),
                    matched_cell=matched_cell,
                    proposed_update_id=update_id,
                )
                continue

            year = int(candidate["year"])
            old_value = get_current_dashboard_value(mapping["indicator"], mapping.get("series_code"), year)
            new_value = float(candidate["new_value"])
            difference = None if old_value is None else new_value - old_value
            update_id = next_proposed_update_id(report_id)
            result_status = "extracted"
            reason = "extracted"
            proposal_status = "Pending Review" if candidate["confidence_score"] < 0.9 else "Ready"
            if candidate["confidence_score"] < 0.8 or "Fuzzy" in str(candidate["extraction_method"]):
                result_status = "ambiguous_match"
                reason = "multiple possible matches found" if candidate["confidence_score"] >= 0.6 else "value empty"
                proposal_status = "Needs Review"

            execute(
                """
                INSERT INTO proposed_updates (
                    update_id, mapping_id, indicator, series_code, year, old_value, new_value,
                    difference, unit_code, source_report, source_report_id, table_or_sheet, evidence_page,
                    extraction_date, status, reviewer_comment, confidence_score, confidence_label,
                    extraction_method, extraction_note, source_evidence, matched_cell, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    update_id,
                    mapping["mapping_id"],
                    mapping["indicator"],
                    mapping.get("series_code"),
                    year,
                    old_value,
                    new_value,
                    difference,
                    mapping.get("unit_code"),
                    report["report_name"],
                    report_id,
                    candidate["table_or_sheet"],
                    candidate["evidence_page"],
                    utc_now(),
                    proposal_status,
                    None,
                    candidate["confidence_score"],
                    candidate["confidence_label"],
                    candidate["extraction_method"],
                    candidate["extraction_note"],
                    candidate["source_evidence"],
                    candidate["matched_cell"],
                    utc_now(),
                    utc_now(),
                ),
            )
            proposals += 1
            record_result(
                report_id=report_id,
                mapping=mapping,
                status=result_status,
                reason=reason,
                expected_report=expected_report,
                expected_table=expected_table,
                expected_row=expected_row,
                expected_column=expected_column,
                closest_matched_row=expected_row,
                closest_matched_column=expected_column,
                confidence_score=candidate["confidence_score"],
                source_sheet_page=candidate["evidence_page"],
                extracted_value=new_value,
                matched_table=candidate["table_or_sheet"],
                matched_cell=candidate["matched_cell"],
                debug_message=candidate["extraction_note"],
                proposed_update_id=update_id,
            )
        except Exception as exc:
            record_result(
                report_id=report_id,
                mapping=mapping,
                status="error",
                reason="error",
                expected_report=expected_report,
                expected_table=expected_table,
                expected_row=expected_row,
                expected_column=expected_column,
                debug_message=str(exc),
            )

    summary = summarize_report_results(report_id)
    report_status = "Needs Review" if summary["needs_review"] or summary["errors"] else "Processed"
    execute(
        "UPDATE reports SET status = ?, extraction_summary = ? WHERE report_id = ?",
        (
            report_status,
            f"Checked {summary['total_mapping_rows_checked']} mapping row(s): {summary['extracted']} extracted, {summary['needs_review']} needs review, {summary['not_found']} not found, {summary['missing_mapping']} missing mapping, {summary['errors']} errors.",
            report_id,
        ),
    )
    log_action(
        "system",
        "report_extracted",
        "report",
        report_id,
        new_value=proposals,
        source_file=report["original_file_name"],
        details={"proposals": proposals, "sheets": [sheet.name for sheet in snapshots], **summary},
    )

    return {
        "report_id": report_id,
        "report_status": report_status,
        "extracted_tables": len(snapshots),
        "proposed_updates": proposals,
        "total_mapping_rows_checked": summary["total_mapping_rows_checked"],
        "extracted_values": summary["extracted"],
        "needs_review": summary["needs_review"],
        "not_found": summary["not_found"],
        "missing_mapping": summary["missing_mapping"],
        "errors": summary["errors"],
        "message": "Excel extraction completed.",
    }
