from __future__ import annotations

"""Mapping-workbook import and approved-dashboard export helpers.

Imports the corrected NISR mapping workbook into the automation database and
exports approved updates into a generated copy of the public SDG workbook.
"""

import re
import shutil
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..database import CONTROL_WORKBOOK_PATH, DATA_DIR, EXPORTS_DIR, execute, fetch_all, fetch_one, get_connection, init_db, init_storage, table_count
from .automation_rules import (
    OBSERVATION_DIMENSION_FIELDS,
    dashboard_year,
    dimension_summary,
    dimension_values,
    identity_where_clause,
    observation_identity,
)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def generate_id(prefix: str, index: int) -> str:
    return f"{prefix}-{index:03d}"


def to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def to_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def parse_table_reference(reference: Any) -> tuple[str | None, str | None]:
    raw_value = str(reference or "").strip()
    if not raw_value:
        return None, None

    compact_value = " ".join(raw_value.replace("\n", " ").split())
    match = re.match(r"^(table|figure)\s+([a-z0-9]+(?:\s*\.\s*[a-z0-9]+)*)\s*[:.-]?\s*(.*)$", compact_value, re.IGNORECASE)
    if not match:
        return None, compact_value

    item_type, number, title = match.groups()
    normalized_number = re.sub(r"\s*\.\s*", ".", number.strip())
    table_number = f"{item_type.title()} {normalized_number}"
    clean_title = title.strip(" :.-") or compact_value
    return table_number, clean_title


def infer_report_family(data_source: str) -> str | None:
    normalized = (data_source or "").lower()
    if "eicv" in normalized:
        return "EICV"
    if "rphc" in normalized or "census" in normalized:
        return "Census"
    if "dhs" in normalized:
        return "DHS"
    return None


def infer_row_label(ref_area: str, province: str, district: str, urbanization: str) -> str | None:
    if district:
        return district
    if province:
        return province
    if urbanization:
        return urbanization
    if ref_area.upper() in {"RW", "RWA", "Rwanda".upper()}:
        return "Rwanda"
    return None


def _data_row_count(workbook_path: Path) -> int:
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "Data" not in workbook.sheetnames:
            return 0
        sheet = workbook["Data"]
        return max(int(sheet.max_row or 0) - 1, 0)
    except Exception:
        return 0


def _dashboard_data_workbook_path(control_workbook_path: Path) -> Path:
    """Prefer the full public SDG workbook for observation rows when available."""
    candidates = [
        DATA_DIR / "approved" / "2025_RW-SDG_Data.xlsx",
        DATA_DIR / "2025_RW-SDG_Data.xlsx",
        control_workbook_path,
    ]
    existing_candidates = [candidate for candidate in candidates if candidate.exists()]
    if not existing_candidates:
        return control_workbook_path
    return max(existing_candidates, key=_data_row_count)


def _is_primary_dashboard_row(row_payload: dict[str, Any]) -> bool:
    """Primary rows carry indicator-level metadata; disaggregated rows do not."""
    ref_area = str(row_payload.get("Ref_Area") or "").strip().upper()
    if ref_area not in {"RW", "RWA"}:
        return False
    for field in (
        "Province",
        "District",
        "Urbanization",
        "Education ",
        "Education",
        "Occupation ",
        "Occupation",
        "Age",
        "Sex",
    ):
        if str(row_payload.get(field) or "").strip():
            return False
    return True


def workbook_exists() -> bool:
    return CONTROL_WORKBOOK_PATH.exists()


def copy_control_workbook(source_path: str | Path) -> str:
    init_storage()
    source = Path(source_path)
    shutil.copy2(source, CONTROL_WORKBOOK_PATH)
    return str(CONTROL_WORKBOOK_PATH)


def _purge_seeded_reports(connection: Any) -> int:
    """Remove report rows that were auto-seeded from the control workbook."""
    rows = connection.execute(
        """
        SELECT report_id
        FROM reports
        WHERE original_file_name IS NULL
        """
    ).fetchall()
    report_ids = [str(row["report_id"]) for row in rows if row and row["report_id"]]
    if not report_ids:
        return 0

    placeholders = ", ".join("?" for _ in report_ids)
    connection.execute(f"DELETE FROM reports WHERE report_id IN ({placeholders})", report_ids)
    return len(report_ids)


def import_control_workbook(workbook_path: str | Path | None = None, *, force: bool = False) -> dict[str, Any]:
    init_storage()
    init_db()

    workbook_file = Path(workbook_path) if workbook_path else CONTROL_WORKBOOK_PATH
    if not workbook_file.exists():
        raise FileNotFoundError(f"Control workbook not found at {workbook_file}")

    with get_connection() as connection:
        _purge_seeded_reports(connection)

    if table_count("source_mapping") > 0 and not force:
        return {
            "imported": False,
            "workbook_path": str(workbook_file),
            "indicators": table_count("indicators"),
            "mappings": table_count("source_mapping"),
            "dashboard_rows": table_count("dashboard_data"),
            "reports": table_count("reports"),
        }

    workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    now = utc_now()
    mapping_seeds: dict[tuple[str, str], dict[str, Any]] = {}

    with get_connection() as connection:
        cursor = connection.cursor()
        cursor.executescript(
            """
            DELETE FROM approved_updates;
            DELETE FROM proposed_updates;
            DELETE FROM extracted_tables;
            DELETE FROM reports
            WHERE COALESCE(original_file_name, '') = ''
               OR report_type = 'pending';
            DELETE FROM source_mapping;
            DELETE FROM dashboard_data;
            DELETE FROM indicators;
            """
        )

        data_workbook_path = _dashboard_data_workbook_path(workbook_file)
        data_workbook = load_workbook(data_workbook_path, read_only=True, data_only=True)
        data_sheet = data_workbook["Data"]
        data_headers = [cell for cell in next(data_sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        year_columns = [(index, int(header)) for index, header in enumerate(data_headers) if isinstance(header, int)]
        table_ref_index = data_headers.index("Table name and number") if "Table name and number" in data_headers else None

        for row_number, row in enumerate(data_sheet.iter_rows(min_row=2, values_only=True), start=2):
            row_payload = {str(data_headers[position] or "").strip(): row[position] for position in range(len(data_headers))}
            is_primary_row = _is_primary_dashboard_row(row_payload)
            indicator = str(row[0] or "").strip()
            series = str(row[1] or "").strip()
            series_code = str(row[2] or "").strip()
            composite = str(row[3] or "").strip()
            unit_code = str(row[4] or "").strip()
            occupation = str(row[5] or "").strip()
            occupation_code = str(row[6] or "").strip()
            data_source = str(row[14] or "").strip()
            description = str(row[15] or "").strip()
            seats = str(row[16] or "").strip()
            ref_area = str(row[7] or "").strip()
            province = str(row[8] or "").strip()
            district = str(row[9] or "").strip()
            urbanization = str(row[10] or "").strip()
            urbanization_code = str(row[11] or "").strip()
            education = str(row[12] or "").strip()
            education_code = str(row[13] or "").strip()
            age_code = str(row[17] or "").strip()
            age = str(row[18] or "").strip()
            sex_code = str(row[19] or "").strip()
            sex = str(row[20] or "").strip()
            table_reference = str(row[table_ref_index] or "").strip() if table_ref_index is not None else ""
            if not indicator:
                continue

            latest_year = None
            latest_value = None
            for column_index, year in year_columns:
                value = to_float(row[column_index])
                if value is not None:
                    latest_year = year
                    latest_value = value

            existing_indicator = cursor.execute(
                """
                SELECT id
                FROM indicators
                WHERE indicator_code = ?
                  AND COALESCE(series_code, '') = COALESCE(?, '')
                LIMIT 1
                """,
                (indicator, series_code or None),
            ).fetchone()
            if is_primary_row or existing_indicator is None:
                cursor.execute(
                    """
                    INSERT INTO indicators (
                        indicator_code, series, series_code, dashboard_description, unit_code,
                        data_source, latest_year, latest_value, geography, disaggregation,
                        source_table_reference, metadata_json, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(indicator_code, series_code) DO UPDATE SET
                        series = excluded.series,
                        dashboard_description = excluded.dashboard_description,
                        unit_code = excluded.unit_code,
                        data_source = excluded.data_source,
                        latest_year = excluded.latest_year,
                        latest_value = excluded.latest_value,
                        geography = excluded.geography,
                        disaggregation = excluded.disaggregation,
                        source_table_reference = excluded.source_table_reference,
                        metadata_json = excluded.metadata_json,
                        updated_at = excluded.updated_at
                    """,
                    (
                        indicator,
                        series,
                        series_code or None,
                        description,
                        unit_code or None,
                        data_source or None,
                        latest_year,
                        latest_value,
                        ref_area or "RW",
                        dimension_summary(
                            {
                                "ref_area": ref_area,
                                "province": province,
                                "district": district,
                                "urbanization": urbanization,
                                "education": education,
                                "occupation": occupation,
                                "age": age,
                                "sex": sex,
                            }
                        ),
                        table_reference or None,
                        "{}",
                        now,
                        now,
                    ),
                )

            table_no, table_title = parse_table_reference(table_reference)
            mapping_key = (indicator, series_code or "")
            mapping_seed = mapping_seeds.get(mapping_key, {})
            if is_primary_row or not mapping_seed:
                mapping_seeds[mapping_key] = {
                    "indicator": indicator,
                    "series": series,
                    "series_code": series_code or None,
                    "dashboard_description": description or None,
                    "unit_code": unit_code or None,
                    "data_source": data_source or None,
                    "latest_year": latest_year,
                    "latest_value": latest_value,
                    "report_family": infer_report_family(data_source),
                    "report_year": latest_year,
                    "file_type": "pdf+excel",
                    "table_no": table_no,
                    "table_title": table_title,
                    "report_indicator_name": description or series or None,
                    "row_label": infer_row_label(ref_area, province, district, urbanization),
                    "source_table_reference": table_reference or None,
                    "existing": mapping_seed.get("existing", False),
                }

            for column_index, year in year_columns:
                cursor.execute(
                    """
                    INSERT INTO dashboard_data (
                        indicator, series, series_code, composite, unit_code, data_source, description,
                        ref_area, province, district, urbanization, urbanization_code, education, education_code,
                        occupation, occupation_code, age, age_code, sex, sex_code, seats, year,
                        value, table_name_and_number, source_row_number, source_year_column, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        indicator,
                        series,
                        series_code or None,
                        composite or None,
                        unit_code or None,
                        data_source or None,
                        description or None,
                        ref_area or None,
                        province or None,
                        district or None,
                        urbanization or None,
                        urbanization_code or None,
                        education or None,
                        education_code or None,
                        occupation or None,
                        occupation_code or None,
                        age or None,
                        age_code or None,
                        sex or None,
                        sex_code or None,
                        seats or None,
                        year,
                        to_float(row[column_index]),
                        table_reference or None,
                        row_number,
                        str(year),
                        now,
                        now,
                    ),
                )

        mapping_sheet = workbook["NISR_Source_Mapping"]
        mapping_headers = [str(cell or "").strip() for cell in next(mapping_sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        imported_mapping_keys: set[tuple[str, str]] = set()
        for index, row in enumerate(mapping_sheet.iter_rows(min_row=2, values_only=True), start=1):
            payload = {mapping_headers[position]: row[position] for position in range(len(mapping_headers))}
            if not payload.get("Indicator"):
                continue
            mapping_id = str(payload.get("Mapping_ID") or generate_id("MAP", index))
            indicator = str(payload.get("Indicator") or "").strip()
            series_code = str(payload.get("Series_Code") or "").strip()
            seed = mapping_seeds.get((indicator, series_code), {})
            seed_table_no = seed.get("table_no")
            seed_table_title = seed.get("table_title")
            seed_report_family = seed.get("report_family")
            seed_report_indicator_name = seed.get("report_indicator_name")
            seed_row_label = seed.get("row_label")
            cursor.execute(
                """
                INSERT INTO source_mapping (
                    mapping_id, indicator, series, series_code, dashboard_description, unit_code,
                    data_source, latest_year, latest_value, report_family, report_name, report_year,
                    file_type, file_name_or_link, sheet_or_page, source_table_reference, table_no, table_title,
                    report_indicator_name, row_label, column_label, geography, disaggregation,
                    extraction_method, confidence, status, reviewer, notes, mapping_type,
                    source_or_survey_period, publication_year, dashboard_display_year,
                    calculation_or_transformation_rule, required_nisr_review_action, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    mapping_id,
                    indicator,
                    str(payload.get("Series") or "").strip() or None,
                    series_code or None,
                    str(payload.get("Dashboard_Description") or "").strip() or seed.get("dashboard_description"),
                    str(payload.get("Unit_Code") or "").strip() or seed.get("unit_code"),
                    str(payload.get("Data_Source") or "").strip() or seed.get("data_source"),
                    to_int(payload.get("Latest_Year")) or seed.get("latest_year"),
                    to_float(payload.get("Latest_Value")) if payload.get("Latest_Value") not in (None, "") else seed.get("latest_value"),
                    str(payload.get("Report_Family") or "").strip() or seed_report_family,
                    str(payload.get("Report_Name") or "").strip() or None,
                    to_int(payload.get("Report_Year")) or seed.get("report_year"),
                    str(payload.get("File_Type") or "").strip() or seed.get("file_type"),
                    str(payload.get("File_Name_or_Link") or "").strip() or None,
                    str(payload.get("Sheet_or_Page") or "").strip() or None,
                    seed.get("source_table_reference"),
                    str(payload.get("Table_No") or "").strip() or seed_table_no,
                    str(payload.get("Table_Title") or "").strip() or seed_table_title,
                    str(payload.get("Report_Indicator_Name") or "").strip() or seed_report_indicator_name,
                    str(payload.get("Row_Label") or "").strip() or seed_row_label,
                    str(payload.get("Column_Label") or "").strip() or None,
                    str(payload.get("Geography") or "").strip() or None,
                    str(payload.get("Disaggregation") or "").strip() or None,
                    str(payload.get("Extraction_Method") or "").strip() or None,
                    str(payload.get("Confidence") or "").strip() or None,
                    str(payload.get("Status") or "").strip() or None,
                    str(payload.get("Reviewer") or "").strip() or None,
                    str(payload.get("Notes") or "").strip() or None,
                    str(payload.get("Mapping_Type") or "").strip() or None,
                    str(payload.get("Source_or_Survey_Period") or "").strip() or None,
                    to_int(payload.get("Publication_Year")),
                    to_int(payload.get("Dashboard_Display_Year")),
                    str(payload.get("Calculation_or_Transformation_Rule") or "").strip() or None,
                    str(payload.get("Required_NISR_Review_Action") or "").strip() or None,
                    now,
                    now,
                ),
            )
            dimension_column_aliases = {
                "ref_area": ("Ref_Area", "Reference_Area", "Geography_Code"),
                "province": ("Province",),
                "district": ("District",),
                "urbanization": ("Urbanization",),
                "urbanization_code": ("Urbanization_Code",),
                "education": ("Education", "Education "),
                "education_code": ("Education_Code",),
                "occupation": ("Occupation", "Occupation "),
                "occupation_code": ("Occupation_Code",),
                "composite": ("Composite", "Composite "),
                "age": ("Age", "Age_Group"),
                "age_code": ("Age_Code",),
                "sex": ("Sex",),
                "sex_code": ("Sex_Code",),
            }
            dimension_values_for_mapping: dict[str, str | None] = {}
            for field, aliases in dimension_column_aliases.items():
                value = next((str(payload.get(alias) or "").strip() for alias in aliases if str(payload.get(alias) or "").strip()), "")
                dimension_values_for_mapping[field] = value or seed.get(field)
            cursor.execute(
                """
                UPDATE source_mapping
                SET ref_area = ?, province = ?, district = ?, urbanization = ?, urbanization_code = ?,
                    education = ?, education_code = ?, occupation = ?, occupation_code = ?, composite = ?,
                    age = ?, age_code = ?, sex = ?, sex_code = ?
                WHERE mapping_id = ?
                """,
                (
                    dimension_values_for_mapping.get("ref_area"),
                    dimension_values_for_mapping.get("province"),
                    dimension_values_for_mapping.get("district"),
                    dimension_values_for_mapping.get("urbanization"),
                    dimension_values_for_mapping.get("urbanization_code"),
                    dimension_values_for_mapping.get("education"),
                    dimension_values_for_mapping.get("education_code"),
                    dimension_values_for_mapping.get("occupation"),
                    dimension_values_for_mapping.get("occupation_code"),
                    dimension_values_for_mapping.get("composite"),
                    dimension_values_for_mapping.get("age"),
                    dimension_values_for_mapping.get("age_code"),
                    dimension_values_for_mapping.get("sex"),
                    dimension_values_for_mapping.get("sex_code"),
                    mapping_id,
                ),
            )
            imported_mapping_keys.add((indicator, series_code))

        # Do not generate AUTO-* mappings. The corrected NISR_Source_Mapping sheet is
        # the source of truth; unmapped indicators must stay in manual review.

        connection.commit()

    return {
        "imported": True,
        "workbook_path": str(workbook_file),
        "indicators": table_count("indicators"),
        "mappings": table_count("source_mapping"),
        "dashboard_rows": table_count("dashboard_data"),
        "reports": table_count("reports"),
    }


def export_updated_dashboard_workbook() -> Path:
    if not CONTROL_WORKBOOK_PATH.exists():
        raise FileNotFoundError("Control workbook is missing.")

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_path = EXPORTS_DIR / f"updated_sdg_dashboard_{timestamp}.xlsx"
    shutil.copy2(CONTROL_WORKBOOK_PATH, export_path)

    workbook = load_workbook(export_path)
    sheet = workbook["Data"]
    approved_rows = fetch_all("SELECT * FROM approved_updates ORDER BY approved_at ASC, id ASC")

    for approved in approved_rows:
        mapping = fetch_one("SELECT * FROM source_mapping WHERE mapping_id = ?", (approved["mapping_id"],))
        if not mapping:
            continue
        target_year = int(approved.get("dashboard_year") or approved.get("year") or dashboard_year(mapping) or 0)
        if target_year <= 0:
            continue
        column_number = ensure_year_column(workbook, sheet, target_year)
        dashboard_row = find_dashboard_observation(
            {**mapping, **{field: approved.get(field) for field in OBSERVATION_DIMENSION_FIELDS}},
            target_year,
        )
        row_number = dashboard_row.get("source_row_number") if dashboard_row else None
        if row_number:
            sheet.cell(row=int(row_number), column=column_number).value = approved["new_value"]

    workbook.save(export_path)
    return export_path


def _year_headers(sheet: Any) -> dict[int, int]:
    lookup: dict[int, int] = {}
    for index, cell in enumerate(sheet[1], start=1):
        value = cell.value
        if isinstance(value, int):
            lookup[value] = index
        elif isinstance(value, str) and value.strip().isdigit():
            lookup[int(value.strip())] = index
    return lookup


def ensure_year_column(workbook: Any, sheet: Any, year: int) -> int:
    year_lookup = _year_headers(sheet)
    if year in year_lookup:
        update_num_columns(workbook, len(year_lookup))
        return year_lookup[year]

    headers = [cell.value for cell in sheet[1]]
    table_column = next((index + 1 for index, value in enumerate(headers) if str(value or "").strip() == "Table name and number"), sheet.max_column + 1)
    insert_at = table_column
    previous_column = max(year_lookup.values(), default=insert_at - 1)
    sheet.insert_cols(insert_at)
    if previous_column > 0:
        for row in range(1, sheet.max_row + 1):
            source_cell = sheet.cell(row=row, column=previous_column)
            target_cell = sheet.cell(row=row, column=insert_at)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.number_format = source_cell.number_format
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
    sheet.cell(row=1, column=insert_at).value = year
    sheet.column_dimensions[get_column_letter(insert_at)].width = sheet.column_dimensions[get_column_letter(previous_column)].width if previous_column > 0 else 12
    update_num_columns(workbook, len(_year_headers(sheet)))
    return insert_at


def update_num_columns(workbook: Any, year_count: int) -> None:
    if "Parameters" not in workbook.sheetnames:
        return
    sheet = workbook["Parameters"]
    for row in sheet.iter_rows():
        for cell in row:
            if str(cell.value or "").strip() == "NumColumns":
                sheet.cell(row=cell.row, column=cell.column + 1).value = year_count
                return


def find_dashboard_observation(mapping: dict[str, Any], year: int) -> dict[str, Any] | None:
    has_explicit_dimensions = any(field in mapping for field in OBSERVATION_DIMENSION_FIELDS)
    if has_explicit_dimensions:
        identity = {
            "indicator": mapping["indicator"],
            "series_code": mapping.get("series_code") or "",
            "year": year,
        }
        for field in OBSERVATION_DIMENSION_FIELDS:
            identity[field] = mapping.get(field) or ""
        where_clause, params = identity_where_clause(identity)
        exact_row = fetch_one(f"SELECT * FROM dashboard_data WHERE {where_clause} LIMIT 1", params)
        if exact_row:
            return exact_row

    candidates = fetch_all(
        """
        SELECT *
        FROM dashboard_data
        WHERE indicator = ?
          AND COALESCE(series_code, '') = COALESCE(?, '')
          AND year = ?
        ORDER BY CASE WHEN ref_area = 'RW' THEN 0 ELSE 1 END, id
        """,
        (mapping["indicator"], mapping.get("series_code"), year),
    )
    if not candidates:
        candidates = fetch_all(
            """
            SELECT *
            FROM dashboard_data
            WHERE indicator = ?
              AND COALESCE(series_code, '') = COALESCE(?, '')
            ORDER BY CASE WHEN ref_area = 'RW' THEN 0 ELSE 1 END, year DESC, id
            """,
            (mapping["indicator"], mapping.get("series_code")),
        )
    if not candidates:
        return None

    def score_candidate(row: dict[str, Any]) -> int:
        score = 0
        mapping_labels = {
            str(mapping.get("row_label") or "").strip().lower(),
            str(mapping.get("geography") or "").strip().lower(),
            str(mapping.get("disaggregation") or "").strip().lower(),
        }
        mapping_labels.discard("")
        for field in (
            "ref_area",
            "province",
            "district",
            "urbanization",
            "education",
            "occupation",
            "composite",
            "age",
            "sex",
        ):
            value = str(row.get(field) or "").strip().lower()
            if value and value in mapping_labels:
                score += 3
        for field in (
            "urbanization_code",
            "education_code",
            "occupation_code",
            "age_code",
            "sex_code",
            "composite",
        ):
            value = str(row.get(field) or "").strip().lower()
            if value and value in mapping_labels:
                score += 2
        if str(row.get("ref_area") or "").upper() == "RW":
            score += 1
        return score

    best_row = max(candidates, key=score_candidate)
    identity = observation_identity(best_row, year=year)
    where_clause, params = identity_where_clause(identity)
    row = fetch_one(f"SELECT * FROM dashboard_data WHERE {where_clause} LIMIT 1", params)
    return row or best_row


def get_current_dashboard_value(
    indicator: str,
    series_code: str | None,
    year: int,
    target: dict[str, Any] | None = None,
) -> float | None:
    if target:
        row = find_dashboard_observation({**target, "indicator": indicator, "series_code": series_code}, year)
        return float(row["value"]) if row and row["value"] is not None else None

    row = fetch_one(
        """
        SELECT value
        FROM dashboard_data
        WHERE indicator = ?
          AND COALESCE(series_code, '') = COALESCE(?, '')
          AND year = ?
        ORDER BY CASE WHEN ref_area = 'RW' THEN 0 ELSE 1 END, id
        LIMIT 1
        """,
        (indicator, series_code, year),
    )
    return float(row["value"]) if row and row["value"] is not None else None
