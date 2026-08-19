from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import database
from app.database import execute, fetch_all, fetch_one, init_db
from app.routes.approvals import review_update
from app.routes.extraction_results import list_extraction_results
from app.routes.proposed_updates import update_proposed_update
from app.schemas import ProposedUpdateEditRequest, ReviewActionRequest
from app.services.automation_rules import (
    identity_where_clause,
    mapping_eligibility,
    observation_identity,
    validate_candidate,
)
from app.services.extraction_results import create_review_placeholder
from app.services.excel_extractor import SheetSnapshot, _dashboard_observation_targets, _direct_series_year_match, _report_matches_mapping
from app.services.excel_extractor import _mapping_reference_candidates, _sheet_reference_score
from app.services.report_cleanup import delete_report
from app.services.workbook_importer import ensure_year_column, export_updated_dashboard_workbook, find_dashboard_observation


@pytest.fixture()
def temp_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    db_path = tmp_path / "nisr_test.db"
    monkeypatch.setattr(database, "DATABASE_URL", "")
    monkeypatch.setattr(database, "DB_PATH", db_path)
    init_db()
    return db_path


def _ready_mapping(**overrides: object) -> dict[str, object]:
    mapping: dict[str, object] = {
        "mapping_id": "MAP-001",
        "indicator": "1.2.1",
        "series_code": "SI_POV_NAHC",
        "unit_code": "PT",
        "status": "Ready",
        "mapping_type": "Direct",
        "dashboard_display_year": 2024,
        "report_year": 2025,
        "publication_year": 2025,
        "table_no": "Table 9.1",
        "table_title": "Headcount poverty rate",
        "row_label": "Rwanda",
        "column_label": "Total",
    }
    mapping.update(overrides)
    return mapping


def test_ready_mapping_eligibility_and_non_ready_blocking() -> None:
    assert mapping_eligibility(_ready_mapping()).eligible

    blocked_statuses = ["Needs Review", "Proxy—Approval Required", "To Map", "Unresolved", "Auto-mapped"]
    for status in blocked_statuses:
        decision = mapping_eligibility(_ready_mapping(status=status))
        assert not decision.eligible
        assert decision.status in {"ambiguous_match", "missing_mapping"}

    auto_decision = mapping_eligibility(_ready_mapping(mapping_id="AUTO-001", status="Ready"))
    assert not auto_decision.eligible
    assert "Auto-generated" in auto_decision.reason


def test_mpi_preservation_blocks_wrong_value_and_keeps_expected_value() -> None:
    mapping = _ready_mapping(
        mapping_id="MAP-002",
        indicator="1.2.2",
        series_code="SI_POV_NMPI",
        unit_code="IX",
        latest_value=0.136,
        calculation_or_transformation_rule="Direct extract: National MPI (M0) index from Table B.5.",
    )
    assert validate_candidate(mapping=mapping, new_value=0.136, old_value=0.136) == []

    warnings = validate_candidate(mapping=mapping, new_value=27.4, old_value=0.136)
    assert any("MPI preservation" in warning for warning in warnings)


def test_mpi_review_placeholder_keeps_candidate_value(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, latest_value, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-MPI", "1.2.2", "SI_POV_NMPI", "IX", 0.136, "Ready", "Direct", 2024, "Table B.5", "National", "MPI (M0)", now, now),
    )
    execute("INSERT INTO reports (report_id, report_name, report_type, upload_date, status) VALUES (?, ?, ?, ?, ?)", ("REP-MPI", "MPI Report", "pdf", now, "Uploaded"))
    mapping = fetch_one("SELECT * FROM source_mapping WHERE mapping_id = ?", ("MAP-MPI",))
    report = fetch_one("SELECT * FROM reports WHERE report_id = ?", ("REP-MPI",))

    update_id = create_review_placeholder(
        report_id="REP-MPI",
        report=report,
        mapping=mapping,
        year=2024,
        old_value=0.129,
        status="Needs Review",
        confidence_score=0.95,
        confidence_label="High",
        extraction_method="PDF ambiguous page match",
        extraction_note="Multiple possible matches found.",
        table_or_sheet="Table B.5",
        evidence_page="42",
        source_evidence="Page 42",
    )
    update = fetch_one("SELECT new_value, difference, validation_warnings, extraction_note FROM proposed_updates WHERE update_id = ?", (update_id,))

    assert update["new_value"] == 0.136
    assert round(update["difference"], 3) == 0.007
    assert "Value empty" not in update["validation_warnings"]
    assert "Review candidate value set from corrected MPI mapping: 0.136" in update["extraction_note"]


def test_review_placeholder_keeps_latest_candidate_value_for_needs_review(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, latest_year, latest_value,
            status, mapping_type, dashboard_display_year, table_no, row_label, column_label,
            created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-REVIEW-VALUE", "1.2.1", "SI_POV_NAHC", "PT", 2024, 27.4, "Ready", "Direct", 2024, "Table 5.1", "Rwanda", "2024", now, now),
    )
    execute("INSERT INTO reports (report_id, report_name, report_type, upload_date, status) VALUES (?, ?, ?, ?, ?)", ("REP-REVIEW-VALUE", "Poverty Report", "excel", now, "Uploaded"))
    mapping = fetch_one("SELECT * FROM source_mapping WHERE mapping_id = ?", ("MAP-REVIEW-VALUE",))
    report = fetch_one("SELECT * FROM reports WHERE report_id = ?", ("REP-REVIEW-VALUE",))

    update_id = create_review_placeholder(
        report_id="REP-REVIEW-VALUE",
        report=report,
        mapping=mapping,
        year=2024,
        old_value=38.2,
        status="Needs Review",
        confidence_score=0.8,
        confidence_label="Medium",
        extraction_method="Excel ambiguous sheet match",
        extraction_note="Multiple possible matches found.",
        table_or_sheet="Table 5.1",
        evidence_page="Table 5.1",
        source_evidence="Table 5.1!B6",
    )
    update = fetch_one("SELECT new_value, validation_warnings, extraction_note FROM proposed_updates WHERE update_id = ?", (update_id,))

    assert update["new_value"] == 27.4
    assert "Value empty" not in update["validation_warnings"]
    assert "Review candidate value set from corrected mapping latest value: 27.4" in update["extraction_note"]


def test_direct_extraction_uses_dashboard_year_not_publication_year() -> None:
    sheet = SheetSnapshot(
        name="Data",
        rows=[
            ["Indicator", "Series_Code", 2024],
            ["1.2.1", "SI_POV_NAHC", 27.4],
        ],
        header_index=0,
        header=["Indicator", "Series_Code", 2024],
        header_lookup={},
        text_blob="Indicator | Series_Code | 2024\n1.2.1 | SI_POV_NAHC | 27.4",
    )
    candidate = _direct_series_year_match(sheet, _ready_mapping(report_year=2025, dashboard_display_year=2024), {"publication_year": 2025})
    assert candidate is not None
    assert candidate["year"] == 2024
    assert candidate["new_value"] == 27.4


def test_direct_extraction_matches_sex_age_dimension_row() -> None:
    sheet = SheetSnapshot(
        name="Data",
        rows=[
            ["Indicator", "Series_Code", "Ref_Area", "Age", "Sex", 2024],
            ["1.2.1", "SI_POV_NAHC", "RW", None, None, 27.4],
            ["1.2.1", "SI_POV_NAHC", "RW", "16+", "Male", 24.9],
            ["1.2.1", "SI_POV_NAHC", "RW", "16+", "Female", 25.4],
        ],
        header_index=0,
        header=["Indicator", "Series_Code", "Ref_Area", "Age", "Sex", 2024],
        header_lookup={},
        text_blob="Indicator | Series_Code | Ref_Area | Age | Sex | 2024",
    )
    mapping = _ready_mapping(ref_area="RW", age="16+", sex="Female", _dimension_target=True)

    candidate = _direct_series_year_match(sheet, mapping, {"publication_year": 2025})

    assert candidate is not None
    assert candidate["new_value"] == 25.4
    assert "Age: 16+ | Sex: Female" in candidate["extraction_note"]


def test_ready_mapping_expands_to_dashboard_observation_targets(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO dashboard_data (
            indicator, series_code, unit_code, ref_area, age, sex, year, value,
            source_row_number, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("1.2.1", "SI_POV_NAHC", "PT", "RW", "16+", "Male", 2024, 24.9, 4, now, now),
    )
    execute(
        """
        INSERT INTO dashboard_data (
            indicator, series_code, unit_code, ref_area, age, sex, year, value,
            source_row_number, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("1.2.1", "SI_POV_NAHC", "PT", "RW", "16+", "Female", 2024, 25.4, 5, now, now),
    )

    targets = _dashboard_observation_targets(_ready_mapping(), 2024)

    assert [target["_dimension_summary"] for target in targets] == ["Age: 16+ | Sex: Male", "Age: 16+ | Sex: Female"]
    assert all(target["_dimension_target"] for target in targets)


def test_dashboard_source_table_reference_matches_alternate_excel_table() -> None:
    sheet = SheetSnapshot(
        name="Table 5.1.",
        rows=[
            ["Table 5.1. Headcount Poverty Rate in 2024 (actual) and 2017 (modelled) by area and province"],
            ["", "", "Total Poverty"],
        ],
        header_index=1,
        header=["", "", "Total Poverty"],
        header_lookup={},
        text_blob="Table 5.1. Headcount Poverty Rate in 2024 (actual) and 2017 (modelled) by area and province",
    )
    references = _mapping_reference_candidates(
        _ready_mapping(table_no="Table 9.1", table_title="Headcount poverty rate (%)—EICV7 (2023/24)"),
        {
            "source_table_reference": "Table 5.1, Headcount poverty rate (%)—EICV7 (2023/24), p. 79, EICV7 Main Indicators Report 2023/24 (2025)"
        },
    )

    assert "Table 5.1" in references
    score, matched_reference = _sheet_reference_score(sheet, references)
    assert score >= 0.9
    assert matched_reference == "Table 5.1"


def test_excel_companion_report_can_process_pdf_mapped_source() -> None:
    mapping = _ready_mapping(file_type="PDF", report_family="EICV", report_name="EICV7 Main Indicators Report 2023/24")
    report = {
        "report_name": "EICV7_Tables_Rwanda_Poverty_Profile",
        "report_family": "EICV",
        "report_type": "excel",
    }

    assert _report_matches_mapping(report, mapping)


def test_validation_detects_duplicates_invalid_percentages_and_unit_mismatch() -> None:
    mapping = _ready_mapping(unit_code="PT")
    warnings = validate_candidate(
        mapping=mapping,
        new_value=125,
        old_value=50,
        existing_unit="IX",
        duplicate=True,
        approved_duplicate=True,
    )
    assert "Percentage outside 0–100" in warnings
    assert "Duplicate observation" in warnings
    assert "Attempted overwrite of approved data" in warnings
    assert any("Definition or unit mismatch" in warning for warning in warnings)


def test_dimension_aware_matching_prefers_matching_disaggregation(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-SEX", "5.1.1", "SERIES", "PT", "Ready", "Direct", 2024, "Table 1", "Female", "Total", now, now),
    )
    for sex, value in [("Male", 10), ("Female", 20)]:
        execute(
            """
            INSERT INTO dashboard_data (
                indicator, series_code, unit_code, ref_area, sex, sex_code, year, value,
                source_row_number, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("5.1.1", "SERIES", "PT", "RW", sex, sex.upper(), 2024, value, 2 if sex == "Male" else 3, now, now),
        )

    mapping = fetch_one("SELECT * FROM source_mapping WHERE mapping_id = ?", ("MAP-SEX",))
    assert mapping is not None
    matched = find_dashboard_observation(mapping, 2024)
    assert matched is not None
    assert matched["sex"] == "Female"
    identity = observation_identity(matched, 2024)
    clause, params = identity_where_clause(identity)
    assert "series_code" in clause
    assert 2024 in params


def test_approval_correction_rejection_return_and_audit(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-APP", "1.2.1", "SI_POV_NAHC", "PT", "Ready", "Direct", 2024, "Table 9.1", "Rwanda", "Total", now, now),
    )
    execute(
        """
        INSERT INTO dashboard_data (
            indicator, series_code, unit_code, ref_area, year, value, source_row_number, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("1.2.1", "SI_POV_NAHC", "PT", "RW", 2024, 25.0, 2, now, now),
    )
    execute(
        """
        INSERT INTO reports (
            report_id, report_name, report_type, upload_date, status
        ) VALUES (?, ?, ?, ?, ?)
        """,
        ("REP", "Report", "excel", now, "Uploaded"),
    )
    for suffix, value in [("APPROVE", 27.4), ("CORRECT", 26.0), ("REJECT", 28.0), ("RETURN", 29.0)]:
        year = 2025 if suffix == "CORRECT" else 2024
        old_value = None if suffix == "CORRECT" else 25.0
        execute(
            """
            INSERT INTO proposed_updates (
                update_id, mapping_id, indicator, series_code, year, old_value, new_value,
                difference, unit_code, source_report, source_report_id, table_or_sheet, evidence_page,
                extraction_date, status, confidence_score, confidence_label, extraction_method,
                extraction_note, source_evidence, matched_cell, dashboard_year, mapping_status,
                mapping_type, validation_warnings, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"UPD-{suffix}",
                "MAP-APP",
                "1.2.1",
                "SI_POV_NAHC",
                year,
                old_value,
                value,
                None if old_value is None else value - old_value,
                "PT",
                "Report",
                "REP",
                "Table",
                "1",
                now,
                "Ready",
                0.95,
                "High",
                "test",
                "test",
                "Sheet!A1",
                "A1",
                year,
                "Ready",
                "Direct",
                "[]",
                now,
                now,
            ),
        )

    review_update("UPD-APPROVE", ReviewActionRequest(action="approve", reviewer="tester"))
    assert fetch_one("SELECT status FROM proposed_updates WHERE update_id = ?", ("UPD-APPROVE",))["status"] == "Approved"

    update_proposed_update("UPD-CORRECT", ProposedUpdateEditRequest(new_value=26.5, reviewer_comment="corrected"))
    review_update("UPD-CORRECT", ReviewActionRequest(action="correct_approve", reviewer="tester", comment="corrected"))
    assert fetch_one("SELECT status FROM proposed_updates WHERE update_id = ?", ("UPD-CORRECT",))["status"] == "Corrected and Approved"

    review_update("UPD-REJECT", ReviewActionRequest(action="reject", reviewer="tester"))
    assert fetch_one("SELECT status FROM proposed_updates WHERE update_id = ?", ("UPD-REJECT",))["status"] == "Rejected"

    review_update("UPD-RETURN", ReviewActionRequest(action="return_for_review", reviewer="tester"))
    assert fetch_one("SELECT status FROM proposed_updates WHERE update_id = ?", ("UPD-RETURN",))["status"] == "Needs Review"

    assert fetch_one("SELECT COUNT(*) AS total FROM approved_updates")["total"] == 2
    assert fetch_one("SELECT COUNT(*) AS total FROM sdg_data_version_history")["total"] == 2
    assert fetch_one("SELECT COUNT(*) AS total FROM audit_log WHERE action LIKE 'proposed_update_%'")["total"] >= 4


def test_approval_updates_exact_dimension_observation(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-DIM-APP", "1.2.1", "SI_POV_NAHC", "PT", "Ready", "Direct", 2024, "Table 9.1", "Rwanda", "Poverty rate", now, now),
    )
    for sex, value, row_number in [("Male", 24.9, 4), ("Female", 25.4, 5)]:
        execute(
            """
            INSERT INTO dashboard_data (
                indicator, series_code, unit_code, ref_area, age, sex, year, value,
                source_row_number, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("1.2.1", "SI_POV_NAHC", "PT", "RW", "16+", sex, 2024, value, row_number, now, now),
        )
    execute("INSERT INTO reports (report_id, report_name, report_type, upload_date, status) VALUES (?, ?, ?, ?, ?)", ("REP-DIM-APP", "Report", "excel", now, "Uploaded"))
    execute(
        """
        INSERT INTO proposed_updates (
            update_id, mapping_id, indicator, series_code, year, old_value, new_value,
            difference, unit_code, source_report, source_report_id, table_or_sheet, evidence_page,
            extraction_date, status, confidence_score, confidence_label, extraction_method,
            extraction_note, source_evidence, matched_cell, dashboard_year, ref_area, age, sex,
            mapping_status, mapping_type, validation_warnings, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            "UPD-DIM-APP",
            "MAP-DIM-APP",
            "1.2.1",
            "SI_POV_NAHC",
            2024,
            25.4,
            26.1,
            0.7,
            "PT",
            "Report",
            "REP-DIM-APP",
            "Table 9.1",
            "Sheet",
            now,
            "Ready",
            0.95,
            "High",
            "test",
            "target observation: Age: 16+ | Sex: Female",
            "Sheet!F4",
            "F4",
            2024,
            "RW",
            "16+",
            "Female",
            "Ready",
            "Direct",
            "[]",
            now,
            now,
        ),
    )

    review_update("UPD-DIM-APP", ReviewActionRequest(action="approve", reviewer="tester"))

    assert fetch_one("SELECT value FROM dashboard_data WHERE sex = ?", ("Male",))["value"] == 24.9
    assert fetch_one("SELECT value FROM dashboard_data WHERE sex = ?", ("Female",))["value"] == 26.1
    approved = fetch_one("SELECT age, sex FROM approved_updates WHERE proposed_update_id = ?", ("UPD-DIM-APP",))
    assert approved["age"] == "16+"
    assert approved["sex"] == "Female"


def test_admin_can_approve_overwrite_with_audit_warning(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-OVERRIDE", "1.2.1", "SI_POV_NAHC", "PT", "Ready", "Direct", 2024, "Table 5.1", "Rwanda", "2024", now, now),
    )
    execute(
        """
        INSERT INTO dashboard_data (
            indicator, series_code, unit_code, ref_area, year, value, source_row_number, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("1.2.1", "SI_POV_NAHC", "PT", "RW", 2024, 25.0, 2, now, now),
    )
    execute("INSERT INTO reports (report_id, report_name, report_type, upload_date, status) VALUES (?, ?, ?, ?, ?)", ("REP-OVERRIDE", "Report", "excel", now, "Uploaded"))
    for update_id, new_value in [("UPD-FIRST", 27.4), ("UPD-SECOND", 28.1)]:
        execute(
            """
            INSERT INTO proposed_updates (
                update_id, mapping_id, indicator, series_code, year, old_value, new_value,
                difference, unit_code, source_report, source_report_id, table_or_sheet, evidence_page,
                extraction_date, status, confidence_score, confidence_label, extraction_method,
                extraction_note, source_evidence, matched_cell, dashboard_year, mapping_status,
                mapping_type, validation_warnings, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                update_id,
                "MAP-OVERRIDE",
                "1.2.1",
                "SI_POV_NAHC",
                2024,
                25.0,
                new_value,
                new_value - 25.0,
                "PT",
                "Report",
                "REP-OVERRIDE",
                "Table 5.1",
                "Sheet",
                now,
                "Ready" if update_id == "UPD-FIRST" else "Needs Review",
                0.95,
                "High",
                "test",
                "test",
                "Sheet!B6",
                "B6",
                2024,
                "Ready",
                "Direct",
                "[]",
                now,
                now,
            ),
        )

    review_update("UPD-FIRST", ReviewActionRequest(action="approve", reviewer="tester"))
    review_update("UPD-SECOND", ReviewActionRequest(action="approve", reviewer="tester", comment="Admin confirmed overwrite"))

    assert fetch_one("SELECT value FROM dashboard_data WHERE indicator = ? AND series_code = ? AND year = ?", ("1.2.1", "SI_POV_NAHC", 2024))["value"] == 28.1
    assert fetch_one("SELECT status FROM proposed_updates WHERE update_id = ?", ("UPD-SECOND",))["status"] == "Approved"
    warnings = fetch_one("SELECT validation_warnings FROM approved_updates WHERE proposed_update_id = ?", ("UPD-SECOND",))["validation_warnings"]
    assert "Attempted overwrite of approved data" in warnings
    assert fetch_one("SELECT COUNT(*) AS total FROM sdg_data_version_history WHERE indicator = ?", ("1.2.1",))["total"] == 2


def test_extraction_results_show_current_value_from_proposed_update(temp_db: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-RESULT", "1.2.1", "SI_POV_NAHC", "PT", "Ready", "Direct", 2024, "Table 5.1", "Rwanda", "2024", now, now),
    )
    execute(
        """
        INSERT INTO dashboard_data (
            indicator, series_code, unit_code, ref_area, year, value, source_row_number, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("1.2.1", "SI_POV_NAHC", "PT", "RW", 2023, 25.0, 2, now, now),
    )
    execute("INSERT INTO reports (report_id, report_name, report_type, upload_date, status) VALUES (?, ?, ?, ?, ?)", ("REP-RESULT", "Report", "excel", now, "Uploaded"))
    execute(
        """
        INSERT INTO proposed_updates (
            update_id, mapping_id, indicator, series_code, year, old_value, new_value,
            difference, unit_code, source_report, source_report_id, table_or_sheet, evidence_page,
            extraction_date, status, confidence_score, confidence_label, extraction_method,
            extraction_note, source_evidence, matched_cell, dashboard_year, mapping_status,
            mapping_type, validation_warnings, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            "UPD-RESULT",
            "MAP-RESULT",
            "1.2.1",
            "SI_POV_NAHC",
            2024,
            25.0,
            27.4,
            2.4,
            "PT",
            "Report",
            "REP-RESULT",
            "Table 5.1",
            "Sheet",
            now,
            "Ready",
            0.95,
            "High",
            "test",
            "test",
            "Sheet!B6",
            "B6",
            2024,
            "Ready",
            "Direct",
            "[]",
            now,
            now,
        ),
    )
    execute(
        """
        INSERT INTO extraction_results (
            report_id, mapping_id, indicator, series_code, expected_report, expected_table, expected_row,
            expected_column, status, reason, confidence_score, proposed_update_id, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("REP-RESULT", "MAP-RESULT", "1.2.1", "SI_POV_NAHC", "Report", "Table 5.1", "Rwanda", "2024", "extracted", "extracted", 0.95, "UPD-RESULT", now, now),
    )

    rows = list_extraction_results(report_id="REP-RESULT", status=None)

    assert rows[0].extracted_value == 27.4
    assert rows[0].previous_year == 2023
    assert rows[0].previous_value == 25.0


def test_export_adds_new_year_updates_numcolumns_and_preserves_original(temp_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    workbook_path = tmp_path / "control.xlsx"
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Indicator", "Series", "Series_Code", "Composite", "Unit_Code", "Occupation", "Occupation_Code", "Ref_Area", "Province", "District", "Urbanization", "Urbanization_Code", "Education", "Education_Code", "Data Source", "Description", "Seats", "Age_Code", "Age", "Sex_Code", "Sex", 2023, 2024, "Table name and number"])
    data.append(["1.2.1", "Poverty", "SI_POV_NAHC", "_T", "PT", None, "_T", "RW", None, None, None, "_T", None, "_T", "NISR", "Poverty", None, "_T", None, "_T", None, 25, 27.4, "Table 9.1"])
    params = workbook.create_sheet("Parameters")
    params.append(["Element", "Type", "PosType", "Position", "", "DataStart", "V2"])
    params.append(["FREQ", "DIM", "FIX", "A", None, "NumColumns", 2])
    workbook.save(workbook_path)

    monkeypatch.setattr("app.services.workbook_importer.CONTROL_WORKBOOK_PATH", workbook_path)
    monkeypatch.setattr("app.services.workbook_importer.EXPORTS_DIR", tmp_path / "exports")

    now = "2026-01-01T00:00:00Z"
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, unit_code, status, mapping_type,
            dashboard_display_year, table_no, row_label, column_label, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-EXP", "1.2.1", "SI_POV_NAHC", "PT", "Ready", "Direct", 2025, "Table 9.1", "Rwanda", "Total", now, now),
    )
    execute(
        """
        INSERT INTO dashboard_data (
            indicator, series_code, unit_code, ref_area, year, value, source_row_number, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("1.2.1", "SI_POV_NAHC", "PT", "RW", 2025, None, 2, now, now),
    )
    execute(
        """
        INSERT INTO approved_updates (
            proposed_update_id, mapping_id, indicator, series_code, year, old_value, new_value,
            unit_code, source_report, table_or_sheet, evidence_page, approved_by, approved_at,
            source_evidence, dashboard_year, mapping_status, mapping_type, validation_warnings, approval_action
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("UPD-EXP", "MAP-EXP", "1.2.1", "SI_POV_NAHC", 2025, None, 30.0, "PT", "Report", "Table", "1", "tester", now, "Sheet!A1", 2025, "Ready", "Direct", "[]", "approve"),
    )

    output_path = export_updated_dashboard_workbook()
    original = load_workbook(workbook_path)
    exported = load_workbook(output_path)

    assert 2025 not in [cell.value for cell in original["Data"][1]]
    exported_headers = [cell.value for cell in exported["Data"][1]]
    assert 2025 in exported_headers
    assert exported["Data"].cell(row=2, column=exported_headers.index(2025) + 1).value == 30.0
    assert exported["Data"].cell(row=2, column=exported_headers.index(2023) + 1).value == 25
    assert exported["Parameters"]["G2"].value == 3


def test_ensure_year_column_updates_numcolumns_without_duplicate() -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Indicator", 2023, "Table name and number"])
    params = workbook.create_sheet("Parameters")
    params.append(["Element", "Type", "PosType", "Position", "", "DataStart", "V2"])
    params.append(["FREQ", "DIM", "FIX", "A", None, "NumColumns", 1])

    first_column = ensure_year_column(workbook, data, 2024)
    second_column = ensure_year_column(workbook, data, 2024)
    assert first_column == second_column
    assert params["G2"].value == 2


def test_delete_report_removes_workflow_artifacts_and_file(temp_db: Path, tmp_path: Path) -> None:
    now = "2026-01-01T00:00:00Z"
    uploaded_file = tmp_path / "uploaded.xlsx"
    uploaded_file.write_text("placeholder", encoding="utf-8")
    execute(
        """
        INSERT INTO reports (
            report_id, report_name, report_type, source_institution, publication_year,
            file_path, original_file_name, upload_date, status, extraction_summary, metadata_json
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("REP-DELETE", "Delete Me Report", "excel", "NISR", 2024, str(uploaded_file), "uploaded.xlsx", now, "Uploaded", None, "{}"),
    )
    execute(
        """
        INSERT INTO extracted_tables (
            report_id, sheet_name, page_number, table_number, table_title, header_row_index, preview_json, text_snapshot, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("REP-DELETE", "Sheet1", 1, "Table 1", "Example Table", 0, "[]", "snapshot", now),
    )
    execute(
        """
        INSERT INTO source_mapping (
            mapping_id, indicator, series_code, status, mapping_type, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        ("MAP-DELETE", "1.2.1", "SI_POV_NAHC", "Ready", "Direct", now, now),
    )
    execute(
        """
        INSERT INTO extraction_results (
            report_id, mapping_id, indicator, series_code, status, reason, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("REP-DELETE", "MAP-DELETE", "1.2.1", "SI_POV_NAHC", "extracted", "ok", now, now),
    )
    execute(
        """
        INSERT INTO proposed_updates (
            update_id, mapping_id, indicator, year, source_report, source_report_id, extraction_date,
            status, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("UPD-DELETE", "MAP-DELETE", "1.2.1", 2024, "Delete Me Report", "REP-DELETE", now, "Ready", now, now),
    )
    execute(
        """
        INSERT INTO approved_updates (
            proposed_update_id, mapping_id, indicator, year, source_report, approved_by, approved_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        ("UPD-DELETE", "MAP-DELETE", "1.2.1", 2024, "Delete Me Report", "tester", now),
    )
    execute(
        """
        INSERT INTO sdg_data_version_history (
            proposed_update_id, indicator, year, changed_by, approval_action, source_report, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        ("UPD-DELETE", "1.2.1", 2024, "tester", "approve", "Delete Me Report", now),
    )
    assert uploaded_file.exists()

    result = delete_report("REP-DELETE", "tester")

    assert result["deleted_report_id"] == "REP-DELETE"
    assert not uploaded_file.exists()
    assert fetch_one("SELECT COUNT(*) AS total FROM reports WHERE report_id = ?", ("REP-DELETE",))["total"] == 0
    assert fetch_one("SELECT COUNT(*) AS total FROM extracted_tables WHERE report_id = ?", ("REP-DELETE",))["total"] == 0
    assert fetch_one("SELECT COUNT(*) AS total FROM extraction_results WHERE report_id = ?", ("REP-DELETE",))["total"] == 0
    assert fetch_one("SELECT COUNT(*) AS total FROM proposed_updates WHERE source_report_id = ?", ("REP-DELETE",))["total"] == 0
    assert fetch_one("SELECT COUNT(*) AS total FROM approved_updates WHERE source_report = ?", ("Delete Me Report",))["total"] == 0
    assert fetch_one("SELECT COUNT(*) AS total FROM sdg_data_version_history WHERE source_report = ?", ("Delete Me Report",))["total"] == 0
    assert fetch_one("SELECT COUNT(*) AS total FROM audit_log WHERE action = ?", ("report_deleted",))["total"] == 1
